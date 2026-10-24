import hashlib, json, re, zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from pypdf import PdfReader
from docx import Document
from pptx import Presentation

ROOT = Path(r"C:\Users\DELL\Desktop\scholau\SCHOLAU\imports\level-100-200")
RAW = ROOT / "raw"
OUT = ROOT / "text"
OUT.mkdir(parents=True, exist_ok=True)

NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main",
      "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
      "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}

def sha256(p: Path) -> str:
    h = hashlib.sha256()
    with p.open("rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()

def pdf_text(p: Path) -> str:
    reader = PdfReader(str(p))
    parts = []
    for i, page in enumerate(reader.pages, 1):
        t = page.extract_text() or ""
        parts.append(f"\n\n--- page {i} ---\n{t}")
    return "".join(parts)

def docx_text(p: Path) -> str:
    doc = Document(str(p))
    parts = []
    for para in doc.paragraphs:
        if para.text.strip():
            style = para.style.name if para.style is not None else ""
            parts.append(f"[{style}] {para.text}")
    for ti, table in enumerate(doc.tables, 1):
        parts.append(f"\n--- table {ti} ---")
        for row in table.rows:
            cells = [c.text.strip().replace("\n", " ") for c in row.cells]
            parts.append(" | ".join(cells))
    return "\n".join(parts)

def pptx_text(p: Path) -> str:
    prs = Presentation(str(p))
    parts = []
    for i, slide in enumerate(prs.slides, 1):
        parts.append(f"\n\n--- slide {i} ---")
        title = ""
        if slide.shapes.title is not None and slide.shapes.title.text:
            title = slide.shapes.title.text
            parts.append(f"# {title}")
        for shape in slide.shapes:
            if not getattr(shape, "has_text_frame", False):
                continue
            text = "\n".join(p.text for p in shape.text_frame.paragraphs if p.text.strip())
            if text and text != title:
                parts.append(text)
    return "\n".join(parts)

def html_text(p: Path) -> str:
    raw = p.read_text(encoding="utf-8", errors="replace")
    raw = re.sub(r"(?is)<script.*?>.*?</script>", " ", raw)
    raw = re.sub(r"(?is)<style.*?>.*?</style>", " ", raw)
    raw = re.sub(r"(?i)<br\s*/?>", "\n", raw)
    raw = re.sub(r"(?i)</(p|div|li|h\d|tr)>", "\n", raw)
    raw = re.sub(r"<[^>]+>", " ", raw)
    raw = re.sub(r"[ \t]+", " ", raw)
    raw = re.sub(r"\n{3,}", "\n\n", raw)
    return raw

def binary_strings(p: Path, limit=400_000) -> str:
    data = p.read_bytes()
    # utf-16le runs common in old office
    chars = []
    i = 0
    buf = []
    while i + 1 < len(data) and len("".join(chars)) < limit:
        lo, hi = data[i], data[i+1]
        if hi == 0 and 32 <= lo < 127:
            buf.append(chr(lo))
            i += 2
            continue
        if len(buf) >= 6:
            chars.append("".join(buf))
        buf = []
        i += 1
    if len(buf) >= 6:
        chars.append("".join(buf))
    text = "\n".join(chars)
    if len(text) < 200:
        ascii_runs = re.findall(rb"[\x20-\x7e]{8,}", data)
        text = "\n".join(s.decode("ascii", "ignore") for s in ascii_runs[:4000])
    return text

def extract(p: Path) -> str:
    ext = p.suffix.lower()
    if ext == ".pdf":
        return pdf_text(p)
    if ext == ".docx":
        return docx_text(p)
    if ext == ".pptx":
        return pptx_text(p)
    if ext in {".txt", ".py", ".html", ".htm"}:
        if ext in {".html", ".htm"}:
            return html_text(p)
        return p.read_text(encoding="utf-8", errors="replace")
    if ext in {".doc", ".ppt"}:
        # try openxml misnamed
        if zipfile.is_zipfile(p):
            if ext == ".ppt":
                return pptx_text(p)
            return docx_text(p)
        return binary_strings(p)
    return p.read_text(encoding="utf-8", errors="replace")

files = [p for p in RAW.rglob("*") if p.is_file()]
seen = {}
index = []
for p in files:
    digest = sha256(p)
    rel = str(p.relative_to(RAW))
    if digest in seen:
        index.append({"file": rel, "sha256": digest, "duplicate_of": seen[digest], "chars": 0})
        continue
    seen[digest] = rel
    try:
        text = extract(p)
        err = None
    except Exception as e:
        text = ""
        err = f"{type(e).__name__}: {e}"
    safe = re.sub(r"[^A-Za-z0-9._-]+", "_", rel)[:180]
    out = OUT / f"{safe}.txt"
    out.write_text(text, encoding="utf-8")
    index.append({
        "file": rel,
        "sha256": digest,
        "chars": len(text),
        "text_file": str(out.relative_to(ROOT)),
        "error": err,
        "preview": re.sub(r"\s+", " ", text)[:400],
    })
    print(f"{len(text):7d}  {rel}" + (f"  ERR {err}" if err else ""))

(ROOT / "extract-index.json").write_text(json.dumps(index, indent=2), encoding="utf-8")
print("UNIQUE", len(seen), "FILES", len(files))
